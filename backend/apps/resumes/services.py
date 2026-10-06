"""
apps/resumes/services.py

Text extraction, cleaning, and structured parsing services for Modules 3 & 4.
Supports:
  - PDF files via PyMuPDF (fitz)
  - DOCX files via python-docx (docx)
  - Deterministic parsing via ResumeParser
"""
import re
import fitz  # PyMuPDF
import docx  # python-docx
from django.utils import timezone

from .models import Resume, ParsedResume
from .parser import ResumeParser


class ResumeTextExtractorService:
    """Service to handle text extraction, cleaning, and structured parsing for resumes."""

    @staticmethod
    def extract_pdf_text(file_path):
        """Extract text from a PDF file using PyMuPDF (fitz)."""
        doc = fitz.open(file_path)
        pages_text = []
        for page in doc:
            page_text = page.get_text("text")
            if page_text:
                pages_text.append(page_text)
        doc.close()
        return "\n".join(pages_text)

    @staticmethod
    def extract_docx_text(file_path):
        """Extract text from a DOCX file using python-docx."""
        doc = docx.Document(file_path)
        lines = []

        for p in doc.paragraphs:
            if p.text:
                lines.append(p.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    lines.append(" | ".join(row_text))

        return "\n".join(lines)

    @staticmethod
    def clean_text(raw_text):
        """Clean extracted resume text."""
        if not raw_text:
            return ""

        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', raw_text)
        text = text.replace('\xa0', ' ').replace('\u200b', '')
        text = text.replace('\r\n', '\n').replace('\r', '\n')

        cleaned_lines = []
        for line in text.split('\n'):
            line_clean = re.sub(r'[ \t]+', ' ', line).strip()
            cleaned_lines.append(line_clean)

        result = "\n".join(cleaned_lines)
        result = re.sub(r'\n{3,}', '\n\n', result)
        return result.strip()

    @classmethod
    def parse_and_store_resume(cls, resume: Resume):
        """
        Module 4 – Parse cleaned resume text into structured data and save ParsedResume.
        """
        text_to_parse = resume.cleaned_text or resume.raw_text or ""
        parsed_dict = ResumeParser.parse_resume(text_to_parse)

        parsed_instance, _ = ParsedResume.objects.update_or_create(
            resume=resume,
            defaults={
                'name': parsed_dict.get('name', ''),
                'email': parsed_dict.get('email', ''),
                'phone': parsed_dict.get('phone', ''),
                'linkedin': parsed_dict.get('linkedin', ''),
                'github': parsed_dict.get('github', ''),
                'skills': parsed_dict.get('skills', []),
                'education': parsed_dict.get('education', []),
                'experience': parsed_dict.get('experience', []),
                'projects': parsed_dict.get('projects', []),
                'certifications': parsed_dict.get('certifications', []),
            }
        )
        return parsed_instance

    @classmethod
    def process_resume(cls, resume: Resume):
        """
        Full processing & parsing pipeline for a Resume instance:
        1. Set status to 'processing'
        2. Extract raw text based on file type
        3. Clean extracted text
        4. Set status to 'completed'
        5. Automatically parse structured data (Module 4)
        """
        resume.processing_status = Resume.STATUS_PROCESSING
        resume.processing_error = None
        resume.save(update_fields=['processing_status', 'processing_error'])

        try:
            file_path = resume.file.path

            if resume.file_type == Resume.FILE_TYPE_PDF:
                raw = cls.extract_pdf_text(file_path)
            elif resume.file_type == Resume.FILE_TYPE_DOCX:
                raw = cls.extract_docx_text(file_path)
            else:
                raise ValueError(f"Unsupported file type for extraction: {resume.file_type}")

            cleaned = cls.clean_text(raw)

            resume.raw_text = raw
            resume.cleaned_text = cleaned
            resume.processing_status = Resume.STATUS_COMPLETED
            resume.processed_at = timezone.now()
            resume.processing_error = None
            resume.save(update_fields=[
                'raw_text',
                'cleaned_text',
                'processing_status',
                'processed_at',
                'processing_error',
            ])

            # Auto-trigger structured parsing
            try:
                cls.parse_and_store_resume(resume)
            except Exception:
                pass

            return resume

        except Exception as exc:
            resume.processing_status = Resume.STATUS_FAILED
            resume.processing_error = str(exc)
            resume.save(update_fields=['processing_status', 'processing_error'])
            raise exc
